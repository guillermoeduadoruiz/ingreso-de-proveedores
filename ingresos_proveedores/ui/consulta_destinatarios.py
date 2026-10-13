from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QTableWidget,
    QTableWidgetItem, QPushButton, QHeaderView, QMessageBox,
    QGroupBox, QLineEdit, QDateEdit, QFileDialog, QApplication,
    QAbstractItemView, QInputDialog, QComboBox,
)
from PyQt5.QtCore import Qt, QDate
from PyQt5.QtGui import QCursor
from servicios.ingreso_svc import IngresoService


class ConsultaDestinatariosWidget(QWidget):
    ESTADOS = {1: 'Emitido', 2: 'En Muelle', 3: 'Despachado'}
    COL_NAMES_DEST = ['Destino', 'Nombre', 'Remitos Emitidos', 'Bultos']
    COL_NAMES_REM = ['N° Remito', 'ID Ingreso', 'Fecha Ingreso', 'Fecha Emisión', 'Proveedor', 'Bultos', 'Estado']

    def __init__(self, ingreso_svc: IngresoService):
        super().__init__()
        self._svc = ingreso_svc
        self._filas: list[dict] = []
        self._destinos: list[dict] = []
        self._remitos_ordenados: list[dict] = []
        self._destino_actual: dict | None = None
        self._remito_seleccionado: dict | None = None
        self._sucursal_seleccionada: str | None = None
        self._on_volver_callback = None
        self._setup_ui()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)

        titulo = QLabel('<h2>Consulta por Destinatario</h2>')
        layout.addWidget(titulo)

        filtro_layout = QHBoxLayout()
        filtro_layout.addWidget(QLabel('Desde:'))
        self.dtp_desde = QDateEdit()
        self.dtp_desde.setCalendarPopup(True)
        self.dtp_desde.setDate(QDate.currentDate().addMonths(-1))
        self.dtp_desde.setDisplayFormat('dd/MM/yyyy')
        filtro_layout.addWidget(self.dtp_desde)
        filtro_layout.addWidget(QLabel('Hasta:'))
        self.dtp_hasta = QDateEdit()
        self.dtp_hasta.setCalendarPopup(True)
        self.dtp_hasta.setDate(QDate.currentDate())
        self.dtp_hasta.setDisplayFormat('dd/MM/yyyy')
        filtro_layout.addWidget(self.dtp_hasta)
        self.txt_filtro = QLineEdit()
        self.txt_filtro.setPlaceholderText('Filtrar por destino o nombre...')
        filtro_layout.addWidget(self.txt_filtro, 1)
        self.btn_buscar = QPushButton('Buscar')
        self.btn_buscar.setObjectName('btnPrimary')
        filtro_layout.addWidget(self.btn_buscar)
        self.btn_actualizar = QPushButton('Actualizar')
        filtro_layout.addWidget(self.btn_actualizar)
        layout.addLayout(filtro_layout)

        self.tabla_destinos = QTableWidget()
        self.tabla_destinos.setColumnCount(4)
        self.tabla_destinos.setHorizontalHeaderLabels(self.COL_NAMES_DEST)
        self.tabla_destinos.setSelectionBehavior(QTableWidget.SelectRows)
        self.tabla_destinos.setSelectionMode(QAbstractItemView.SingleSelection)
        self.tabla_destinos.setEditTriggers(QTableWidget.NoEditTriggers)
        self.tabla_destinos.setAlternatingRowColors(True)
        self.tabla_destinos.verticalHeader().setVisible(False)
        self.tabla_destinos.itemSelectionChanged.connect(self._on_destino_seleccionado)
        layout.addWidget(self.tabla_destinos, 1)

        grp_remitos = QGroupBox('Remitos Emitidos')
        det_layout = QVBoxLayout(grp_remitos)
        filtro_rem_layout = QHBoxLayout()
        filtro_rem_layout.addWidget(QLabel('Estado:'))
        self.cmb_estado = QComboBox()
        self.cmb_estado.addItem('Todos los estados', None)
        for cod, desc in self.ESTADOS.items():
            self.cmb_estado.addItem(desc, cod)
        filtro_rem_layout.addWidget(self.cmb_estado)
        self.txt_filtro_remito = QLineEdit()
        self.txt_filtro_remito.setPlaceholderText('Filtrar por N° remito, ID ingreso o proveedor...')
        filtro_rem_layout.addWidget(self.txt_filtro_remito, 1)
        self.btn_limpiar_filtros_rem = QPushButton('Limpiar filtros')
        self.btn_limpiar_filtros_rem.setObjectName('btnSecondary')
        filtro_rem_layout.addWidget(self.btn_limpiar_filtros_rem)
        det_layout.addLayout(filtro_rem_layout)
        self.lbl_remitos_info = QLabel('')
        self.lbl_remitos_info.setStyleSheet('font-size: 12px; color: #2196F3; font-weight: bold;')
        det_layout.addWidget(self.lbl_remitos_info)
        self.tabla_remitos = QTableWidget()
        self.tabla_remitos.setColumnCount(7)
        self.tabla_remitos.setHorizontalHeaderLabels(self.COL_NAMES_REM)
        self.tabla_remitos.setSelectionBehavior(QTableWidget.SelectRows)
        self.tabla_remitos.setEditTriggers(QTableWidget.NoEditTriggers)
        self.tabla_remitos.setAlternatingRowColors(True)
        self.tabla_remitos.verticalHeader().setVisible(False)
        self.tabla_remitos.itemSelectionChanged.connect(self._on_remito_seleccionado)
        det_layout.addWidget(self.tabla_remitos, 1)
        layout.addWidget(grp_remitos, 1)

        grp_acciones = QGroupBox()
        acc_layout = QHBoxLayout(grp_acciones)
        self.btn_volver = QPushButton('← Volver')
        self.btn_volver.setObjectName('btnSecondary')
        acc_layout.addWidget(self.btn_volver)
        self.btn_estado = QPushButton('Actualizar Estado')
        self.btn_estado.setObjectName('btnPrimary')
        self.btn_estado.setEnabled(False)
        acc_layout.addWidget(self.btn_estado)
        acc_layout.addStretch(1)
        self.btn_exportar = QPushButton('Exportar Excel')
        self.btn_exportar.setObjectName('btnPrimary')
        acc_layout.addWidget(self.btn_exportar)
        layout.addWidget(grp_acciones)

        self.btn_buscar.clicked.connect(self._cargar_datos)
        self.btn_actualizar.clicked.connect(self._cargar_datos)
        self.txt_filtro.returnPressed.connect(self._cargar_datos)
        self.cmb_estado.currentIndexChanged.connect(self._poblar_tabla_destinos)
        self.txt_filtro_remito.textChanged.connect(self._poblar_tabla_destinos)
        self.btn_limpiar_filtros_rem.clicked.connect(self._limpiar_filtros_remitos)
        self.btn_volver.clicked.connect(self._volver)
        self.btn_estado.clicked.connect(self._actualizar_estado)
        self.btn_exportar.clicked.connect(self._exportar_excel)

    def set_on_volver(self, callback):
        self._on_volver_callback = callback

    def _volver(self):
        if self._on_volver_callback:
            self._on_volver_callback()

    def cargar_datos(self):
        self._cargar_datos()

    def _cargar_datos(self):
        try:
            desde = self.dtp_desde.date().toPyDate()
            hasta = self.dtp_hasta.date().toPyDate()
            self._filas = self._svc.obtener_consulta_destinatarios(desde, hasta)
            self._poblar_tabla_destinos()
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al cargar datos:\n{e}')

    def _poblar_tabla_destinos(self):
        filtro = self.txt_filtro.text().strip().lower()
        estado = self.cmb_estado.currentData()
        filtro_rem = self.txt_filtro_remito.text().strip().lower()

        filas = self._filas
        if estado is not None:
            filas = [r for r in filas if int(r.get('estado') or 1) == estado]
        if filtro_rem:
            filas = [
                r for r in filas
                if filtro_rem in str(r.get('num_remito') or '').lower()
                or filtro_rem in str(r.get('id_ingreso') or '').lower()
                or filtro_rem in str(r.get('num_proveedor') or '').lower()
            ]

        agrupado: dict[str, dict] = {}
        for r in filas:
            suc = r.get('sucursal_destino') or '(sin destino)'
            if suc not in agrupado:
                agrupado[suc] = {
                    'sucursal': suc,
                    'nombre': r.get('nombre_destino') or '',
                    'remitos': [],
                    'total_bultos': 0,
                }
            agrupado[suc]['remitos'].append(r)
            agrupado[suc]['total_bultos'] += r.get('total_bultos') or 0

        if filtro:
            agrupado = {
                s: d for s, d in agrupado.items()
                if filtro in s.lower() or filtro in d['nombre'].lower()
            }

        self._destinos = sorted(agrupado.values(), key=lambda d: d['sucursal'].lower())

        self.tabla_destinos.setRowCount(len(self._destinos))
        for i, d in enumerate(self._destinos):
            self.tabla_destinos.setItem(i, 0, QTableWidgetItem(d['sucursal']))
            self.tabla_destinos.setItem(i, 1, QTableWidgetItem(d['nombre']))
            item_rem = QTableWidgetItem(str(len(d['remitos'])))
            item_rem.setTextAlignment(Qt.AlignCenter)
            self.tabla_destinos.setItem(i, 2, item_rem)
            item_bult = QTableWidgetItem(str(d['total_bultos']))
            item_bult.setTextAlignment(Qt.AlignCenter)
            self.tabla_destinos.setItem(i, 3, item_bult)

        header = self.tabla_destinos.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.Stretch)

        if self._destinos:
            idx = 0
            if self._sucursal_seleccionada:
                for i, d in enumerate(self._destinos):
                    if d['sucursal'] == self._sucursal_seleccionada:
                        idx = i
                        break
            self.tabla_destinos.selectRow(idx)
        else:
            self.tabla_remitos.setRowCount(0)
            self.lbl_remitos_info.setText('No se encontraron remitos en el período.')

    def _on_destino_seleccionado(self):
        sel = self.tabla_destinos.selectedItems()
        if not sel:
            return
        row = sel[0].row()
        if 0 <= row < len(self._destinos):
            self._poblar_tabla_remitos(self._destinos[row])

    def _on_remito_seleccionado(self):
        sel = self.tabla_remitos.selectedItems()
        self._remito_seleccionado = None
        if sel:
            row = sel[0].row()
            if 0 <= row < len(self._remitos_ordenados):
                self._remito_seleccionado = self._remitos_ordenados[row]
        self.btn_estado.setEnabled(self._remito_seleccionado is not None)

    def _limpiar_filtros_remitos(self):
        self.cmb_estado.setCurrentIndex(0)
        self.txt_filtro_remito.clear()

    def _actualizar_estado(self):
        r = self._remito_seleccionado
        if not r:
            return
        actual = int(r.get('estado') or 1)
        nombres = list(self.ESTADOS.values())
        idx = actual - 1
        if not (0 <= idx < len(nombres)):
            idx = 0
        opcion, ok = QInputDialog.getItem(
            self, 'Actualizar Estado',
            f'Estado del remito {r.get("num_remito") or ""} del ingreso {r.get("id_ingreso") or ""}:',
            nombres, idx, False,
        )
        if not ok:
            return
        nuevo = next((k for k, v in self.ESTADOS.items() if v == opcion), None)
        if nuevo is None or nuevo == actual:
            return
        try:
            self._svc.actualizar_estado_remito_generado(
                r['id_ingreso'], r['sucursal_destino'], nuevo
            )
            self._sucursal_seleccionada = r['sucursal_destino']
            self._cargar_datos()
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al actualizar estado:\n{e}')

    def _poblar_tabla_remitos(self, destino: dict):
        self._destino_actual = destino
        remitos = sorted(destino['remitos'], key=lambda r: str(r.get('num_remito') or ''))
        self._remitos_ordenados = remitos
        self.tabla_remitos.setRowCount(len(remitos))
        for i, r in enumerate(remitos):
            self.tabla_remitos.setItem(i, 0, QTableWidgetItem(str(r.get('num_remito') or '')))
            self.tabla_remitos.setItem(i, 1, QTableWidgetItem(str(r.get('id_ingreso') or '')))
            f_ing = r.get('fecha_ingreso')
            self.tabla_remitos.setItem(i, 2, QTableWidgetItem(f_ing.strftime('%d/%m/%Y') if f_ing else ''))
            f_gen = r.get('fecha_generacion')
            self.tabla_remitos.setItem(i, 3, QTableWidgetItem(f_gen.strftime('%d/%m/%Y %H:%M') if f_gen else ''))
            prov = f"{r.get('num_proveedor') or ''} - {r.get('razon_social') or ''}".strip(' -')
            self.tabla_remitos.setItem(i, 4, QTableWidgetItem(prov))
            item_b = QTableWidgetItem(str(r.get('total_bultos') or 0))
            item_b.setTextAlignment(Qt.AlignCenter)
            self.tabla_remitos.setItem(i, 5, item_b)
            item_e = QTableWidgetItem(self.ESTADOS.get(int(r.get('estado') or 1), 'Emitido'))
            item_e.setTextAlignment(Qt.AlignCenter)
            self.tabla_remitos.setItem(i, 6, item_e)

        header = self.tabla_remitos.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeToContents)
        header.setSectionResizeMode(4, QHeaderView.Stretch)

        info = f'Destinatario: {destino["sucursal"]}'
        if destino['nombre']:
            info += f' - {destino["nombre"]}'
        info += f'  |  {len(remitos)} remito(s) emitidos  |  {destino["total_bultos"]} bultos'
        self.lbl_remitos_info.setText(info)

    def _exportar_excel(self):
        if not self._destinos:
            QMessageBox.information(self, 'Sin datos', 'No hay datos para exportar.')
            return
        path, _ = QFileDialog.getSaveFileName(
            self, 'Exportar a Excel', 'consulta_destinatarios.xlsx',
            'Archivos Excel (*.xlsx)',
        )
        if not path:
            return
        try:
            QApplication.setOverrideCursor(QCursor(Qt.WaitCursor))
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
            wb = Workbook()
            ws = wb.active
            ws.title = 'Destinatarios'

            encabezados = ['Destino', 'Nombre', 'N° Remito', 'ID Ingreso',
                           'Fecha Ingreso', 'Fecha Emisión', 'Proveedor', 'Bultos', 'Estado']
            e_enc = Font(bold=True, color='FFFFFF', size=10)
            r_enc = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')

            e_dst = Font(bold=True, size=10)
            e_rem = Font(size=10)

            for col, h in enumerate(encabezados, 1):
                c = ws.cell(row=1, column=col, value=h)
                c.font = e_enc
                c.fill = r_enc
                c.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 22

            row = 2
            for d in self._destinos:
                remitos = sorted(d['remitos'], key=lambda r: str(r.get('num_remito') or ''))
                for r in remitos:
                    f_ing = r.get('fecha_ingreso')
                    f_gen = r.get('fecha_generacion')
                    ws.cell(row=row, column=1, value=d['sucursal']).font = e_dst
                    ws.cell(row=row, column=2, value=d['nombre']).font = e_dst
                    ws.cell(row=row, column=3, value=str(r.get('num_remito') or '')).font = e_rem
                    ws.cell(row=row, column=4, value=str(r.get('id_ingreso') or '')).font = e_rem
                    ws.cell(row=row, column=5, value=f_ing.strftime('%d/%m/%Y') if f_ing else '').font = e_rem
                    ws.cell(row=row, column=6, value=f_gen.strftime('%d/%m/%Y %H:%M') if f_gen else '').font = e_rem
                    prov = f"{r.get('num_proveedor') or ''} - {r.get('razon_social') or ''}".strip(' -')
                    ws.cell(row=row, column=7, value=prov).font = e_rem
                    ws.cell(row=row, column=8, value=r.get('total_bultos') or 0).font = e_rem
                    ws.cell(row=row, column=8).alignment = Alignment(horizontal='center')
                    ws.cell(row=row, column=9, value=self.ESTADOS.get(int(r.get('estado') or 1), 'Emitido')).font = e_rem
                    ws.cell(row=row, column=9).alignment = Alignment(horizontal='center')
                    row += 1

            for col, dims in [(1, 16), (2, 32), (3, 14), (4, 14), (5, 14), (6, 18), (7, 45), (8, 10), (9, 14)]:
                ws.column_dimensions[get_column_letter(col)].width = dims

            wb.save(path)
            QMessageBox.information(self, 'Exportar', f'Excel generado:\n{path}')
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al exportar:\n{e}')
        finally:
            QApplication.restoreOverrideCursor()
