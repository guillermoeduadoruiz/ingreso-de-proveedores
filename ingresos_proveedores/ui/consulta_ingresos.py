from datetime import date as _date
from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QTableWidget,
    QTableWidgetItem, QPushButton, QHeaderView, QMessageBox,
    QGroupBox, QLineEdit, QInputDialog, QComboBox, QAbstractItemView,
    QCalendarWidget, QDateEdit, QFileDialog, QApplication,
)
from PyQt5.QtCore import Qt, QDate
from PyQt5.QtGui import QFontMetrics, QCursor
from database.conexiones import ConfigSQLServer
from database.sqlserver import SQLServerConnection
from ui.login_dialog import LoginDialog
from servicios.ingreso_svc import IngresoService
from servicios.impresion_svc import ImpresionService


class ConsultaIngresosWidget(QWidget):
    ESTADOS_REMITO = {1: 'Emitido', 2: 'En Muelle', 3: 'Despachado'}
    COL_RECEPCIONES = 5
    COL_NAMES = ['ID Ingreso', 'Fecha', 'Proveedor', 'Razón Social', 'Remitos', 'Remito Generado']
    COL_FILTRO = {0: 'text', 1: 'text', 2: 'text', 3: 'text', 4: 'combo', 5: 'text'}

    def __init__(self, ingreso_svc: IngresoService, clave_anulacion: str = '',
                 sqlserver_config: ConfigSQLServer | None = None,
                 impresion_svc: ImpresionService | None = None):
        super().__init__()
        self._svc = ingreso_svc
        self._clave_anulacion = clave_anulacion
        self._sqlserver_config = sqlserver_config
        self._imp_svc = impresion_svc
        self._cm = impresion_svc._cm if impresion_svc else None
        self._ingresos: list[dict] = []
        self._ingresos_filtrados: list[dict] = []
        self._remitos: list[dict] = []
        self._remito_recepciones: dict[int, list[dict]] = {}
        self._on_volver_callback = None
        self._expanded_rows: set[int] = set()
        self._default_row_heights: dict[int, int] = {}
        self._filtros: dict[int, str] = {0: '', 1: '', 2: '', 3: '', 4: '', 5: ''}
        self._fecha_desde: QDate | None = None
        self._fecha_hasta: QDate | None = None
        self._sort_col: int | None = None
        self._sort_asc: bool = True
        self._remitos_generados_map: dict[str, list[dict]] = {}
        self._setup_ui()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)

        titulo = QLabel('<h2>Consulta de Ingresos Realizados</h2>')
        layout.addWidget(titulo)

        filtro_layout = QHBoxLayout()
        self.txt_filtro = QLineEdit()
        self.txt_filtro.setPlaceholderText('Buscar por ID de Ingreso o Proveedor...')
        filtro_layout.addWidget(self.txt_filtro, 1)
        self.btn_buscar = QPushButton('Buscar')
        self.btn_buscar.setObjectName('btnPrimary')
        filtro_layout.addWidget(self.btn_buscar)
        self.btn_actualizar = QPushButton('Actualizar')
        filtro_layout.addWidget(self.btn_actualizar)
        layout.addLayout(filtro_layout)

        fecha_layout = QHBoxLayout()
        fecha_layout.addWidget(QLabel('Desde:'))
        self.dtp_desde = QDateEdit()
        self.dtp_desde.setCalendarPopup(True)
        self.dtp_desde.setDate(QDate.currentDate().addMonths(-1))
        self.dtp_desde.setDisplayFormat('dd/MM/yyyy')
        self.dtp_desde.dateChanged.connect(self._on_fecha_rango_changed)
        fecha_layout.addWidget(self.dtp_desde)
        fecha_layout.addWidget(QLabel('Hasta:'))
        self.dtp_hasta = QDateEdit()
        self.dtp_hasta.setCalendarPopup(True)
        self.dtp_hasta.setDate(QDate.currentDate())
        self.dtp_hasta.setDisplayFormat('dd/MM/yyyy')
        self.dtp_hasta.dateChanged.connect(self._on_fecha_rango_changed)
        fecha_layout.addWidget(self.dtp_hasta)
        self.btn_limpiar_fechas = QPushButton('Limpiar fechas')
        self.btn_limpiar_fechas.setObjectName('btnSecondary')
        self.btn_limpiar_fechas.clicked.connect(self._limpiar_fechas)
        fecha_layout.addWidget(self.btn_limpiar_fechas)
        fecha_layout.addStretch()
        layout.addLayout(fecha_layout)

        self.tabla_ingresos = QTableWidget()
        self.tabla_ingresos.setColumnCount(6)
        self.tabla_ingresos.setHorizontalHeaderLabels(self.COL_NAMES)
        self.tabla_ingresos.setSelectionBehavior(QTableWidget.SelectRows)
        self.tabla_ingresos.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.tabla_ingresos.setEditTriggers(QTableWidget.NoEditTriggers)
        self.tabla_ingresos.setAlternatingRowColors(True)
        self.tabla_ingresos.verticalHeader().setVisible(False)
        self.tabla_ingresos.horizontalHeader().sectionClicked.connect(self._on_header_clicked)
        layout.addWidget(self.tabla_ingresos, 1)

        grp_detalle = QGroupBox('Detalle del Ingreso')
        det_layout = QVBoxLayout(grp_detalle)

        self.lbl_detalle_id = QLabel('')
        self.lbl_detalle_id.setStyleSheet('font-size: 14px; font-weight: bold; color: #2196F3;')
        det_layout.addWidget(self.lbl_detalle_id)

        self.lbl_detalle_proveedor = QLabel('')
        det_layout.addWidget(self.lbl_detalle_proveedor)

        self.lbl_remitos_generados = QLabel('')
        self.lbl_remitos_generados.setStyleSheet('font-size: 12px; color: #4CAF50; font-weight: bold;')
        det_layout.addWidget(self.lbl_remitos_generados)

        self.tabla_remitos = QTableWidget()
        self.tabla_remitos.setColumnCount(6)
        self.tabla_remitos.setHorizontalHeaderLabels(
            ['N° Remito', 'Fecha', 'Bultos', 'Unidades', 'Valor Declarado', 'Recepciones Asociadas']
        )
        self.tabla_remitos.setSelectionBehavior(QTableWidget.SelectRows)
        self.tabla_remitos.setEditTriggers(QTableWidget.NoEditTriggers)
        self.tabla_remitos.setAlternatingRowColors(True)
        self.tabla_remitos.verticalHeader().setVisible(False)
        det_layout.addWidget(self.tabla_remitos, 1)

        layout.addWidget(grp_detalle, 1)

        grp_acciones = QGroupBox()
        acc_layout = QHBoxLayout(grp_acciones)
        self.btn_volver = QPushButton('← Volver')
        self.btn_volver.setObjectName('btnSecondary')
        acc_layout.addWidget(self.btn_volver)
        acc_layout.addStretch(1)
        self.btn_exportar = QPushButton('Exportar Excel')
        self.btn_exportar.setObjectName('btnPrimary')
        acc_layout.addWidget(self.btn_exportar)
        self.btn_reimprimir_remitos = QPushButton('Reimprimir Remitos')
        self.btn_reimprimir_remitos.setObjectName('btnPrimary')
        self.btn_reimprimir_remitos.setEnabled(False)
        acc_layout.addWidget(self.btn_reimprimir_remitos)
        self.btn_reimprimir_etiquetas = QPushButton('Reimprimir Etiquetas')
        self.btn_reimprimir_etiquetas.setObjectName('btnPrimary')
        self.btn_reimprimir_etiquetas.setEnabled(False)
        acc_layout.addWidget(self.btn_reimprimir_etiquetas)
        self.btn_eliminar = QPushButton('Eliminar Ingreso(s)')
        self.btn_eliminar.setObjectName('btnDanger')
        self.btn_eliminar.setEnabled(False)
        if not self._clave_anulacion:
            self.btn_eliminar.setVisible(False)
        acc_layout.addWidget(self.btn_eliminar)
        layout.addWidget(grp_acciones)

        self.btn_buscar.clicked.connect(self._buscar)
        self.btn_actualizar.clicked.connect(self._cargar_ingresos)
        self.txt_filtro.returnPressed.connect(self._buscar)
        self.tabla_ingresos.itemSelectionChanged.connect(self._on_seleccion_cambiada)
        self.tabla_remitos.cellClicked.connect(self._on_remito_cell_clicked)
        self.btn_volver.clicked.connect(self._volver)
        self.btn_eliminar.clicked.connect(self._eliminar_ingreso)
        self.btn_exportar.clicked.connect(self._exportar_excel)
        self.btn_reimprimir_remitos.clicked.connect(lambda: self._reimprimir('remitos'))
        self.btn_reimprimir_etiquetas.clicked.connect(lambda: self._reimprimir('etiquetas'))

    def set_on_volver(self, callback):
        self._on_volver_callback = callback

    def _volver(self):
        if self._on_volver_callback:
            self._on_volver_callback()

    def cargar_datos(self):
        self._fecha_desde = self.dtp_desde.date()
        self._fecha_hasta = self.dtp_hasta.date()
        self._cargar_ingresos()

    def _cargar_ingresos(self):
        try:
            self._ingresos = self._svc.obtener_ingresos_recientes()
            self._remitos_generados_map = self._svc.obtener_remitos_generados_por_ingresos(
                [i['id_ingreso'] for i in self._ingresos]
            )
            self._aplicar_filtros_y_orden()
            self._limpiar_detalle()
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al cargar ingresos:\n{e}')

    def _buscar(self):
        filtro = self.txt_filtro.text().strip().lower()
        if not filtro:
            self._cargar_ingresos()
            return
        try:
            ingresos = self._svc.obtener_ingresos_recientes(500)
            self._ingresos = [
                i for i in ingresos
                if filtro in i['id_ingreso'].lower()
                or filtro in i.get('razon_social', '').lower()
                or filtro in i.get('num_proveedor', '').lower()
            ]
            self._remitos_generados_map = self._svc.obtener_remitos_generados_por_ingresos(
                [i['id_ingreso'] for i in self._ingresos]
            )
            self._aplicar_filtros_y_orden()
            self._limpiar_detalle()
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al buscar:\n{e}')

    def _on_header_clicked(self, logical_index):
        if logical_index == self._sort_col:
            self._sort_asc = not self._sort_asc
        else:
            self._sort_col = logical_index
            self._sort_asc = True
        self._aplicar_filtros_y_orden()
        self._actualizar_sort_indicators()

    def _actualizar_sort_indicators(self):
        header = self.tabla_ingresos.horizontalHeader()
        header.setSortIndicatorShown(True)
        if self._sort_col is not None:
            order = Qt.AscendingOrder if self._sort_asc else Qt.DescendingOrder
            header.setSortIndicator(self._sort_col, order)
        else:
            header.setSortIndicatorShown(False)

    def _aplicar_filtros_y_orden(self):
        filtrados = self._ingresos[:]

        if self._fecha_desde is not None:
            desde = self._fecha_desde.toPyDate()
            filtrados = [i for i in filtrados if i.get('fecha_ingreso') and i['fecha_ingreso'] >= desde]
        if self._fecha_hasta is not None:
            hasta = self._fecha_hasta.toPyDate()
            filtrados = [i for i in filtrados if i.get('fecha_ingreso') and i['fecha_ingreso'] <= hasta]

        for col, val in self._filtros.items():
            if not val:
                continue
            val_lower = val.lower()
            if col == 0:
                filtrados = [i for i in filtrados if val_lower in i['id_ingreso'].lower()]
            elif col == 1:
                filtrados = [
                    i for i in filtrados
                    if i.get('fecha_ingreso') and val_lower in i['fecha_ingreso'].strftime('%d/%m/%Y').lower()
                ]
            elif col == 2:
                filtrados = [i for i in filtrados if val_lower in i.get('num_proveedor', '').lower()]
            elif col == 3:
                filtrados = [i for i in filtrados if val_lower in i.get('razon_social', '').lower()]
            elif col == 4:
                filtrados = [i for i in filtrados if val == str(i.get('cantidad_remitos', 0))]
            elif col == 5:
                filtrados = [
                    i for i in filtrados
                    if val_lower in self._nums_generados(i['id_ingreso']).lower()
                ]

        if self._sort_col is not None:
            col = self._sort_col
            reverse = not self._sort_asc
            if col == 1:
                filtrados.sort(
                    key=lambda i: i.get('fecha_ingreso') or _date.min,
                    reverse=reverse
                )
            elif col == 4:
                filtrados.sort(key=lambda i: i.get('cantidad_remitos', 0), reverse=reverse)
            elif col == 5:
                filtrados.sort(
                    key=lambda i: self._nums_generados(i['id_ingreso']).lower(),
                    reverse=reverse
                )
            else:
                key_map = {0: 'id_ingreso', 2: 'num_proveedor', 3: 'razon_social'}
                attr = key_map.get(col, 'id_ingreso')
                filtrados.sort(key=lambda i: i.get(attr, '').lower(), reverse=reverse)

        self._ingresos_filtrados = filtrados
        self._poblar_tabla_ingresos()
        self._actualizar_resumen()

    def _nums_generados(self, id_ingreso: str) -> str:
        nums = [
            g.get('num_remito') or ''
            for g in self._remitos_generados_map.get(id_ingreso, [])
        ]
        return ', '.join(n for n in nums if n)

    def _poblar_tabla_ingresos(self):
        self.tabla_ingresos.blockSignals(True)
        datos = self._ingresos_filtrados
        self.tabla_ingresos.setRowCount(len(datos) + 1)

        for col in range(self.tabla_ingresos.columnCount()):
            if col == 0:
                w = QLineEdit()
                w.setPlaceholderText('Filtrar ID...')
                w.setText(self._filtros[0])
                w.setClearButtonEnabled(True)
                w.textChanged.connect(lambda text, c=col: self._on_filtro_changed(c, text))
                self.tabla_ingresos.setCellWidget(0, col, w)
            elif col == 1:
                self.tabla_ingresos.setCellWidget(0, col, self._crear_filtro_fecha())
            elif col == 2:
                self.tabla_ingresos.setCellWidget(0, col, self._crear_combo_filtro(
                    'num_proveedor', 'Filtrar proveedor...'))
            elif col == 3:
                self.tabla_ingresos.setCellWidget(0, col, self._crear_combo_filtro(
                    'razon_social', 'Filtrar razón social...'))
            elif col == 4:
                cbo = QComboBox()
                cbo.addItem('')
                cantidades = sorted({str(i.get('cantidad_remitos', 0)) for i in self._ingresos})
                for v in cantidades:
                    cbo.addItem(v)
                idx = cbo.findText(self._filtros[4])
                if idx >= 0:
                    cbo.setCurrentIndex(idx)
                cbo.currentTextChanged.connect(lambda text, c=4: self._on_filtro_changed(c, text))
                self.tabla_ingresos.setCellWidget(0, col, cbo)
            elif col == 5:
                w = QLineEdit()
                w.setPlaceholderText('Filtrar remito...')
                w.setText(self._filtros[5])
                w.setClearButtonEnabled(True)
                w.textChanged.connect(lambda text, c=5: self._on_filtro_changed(c, text))
                self.tabla_ingresos.setCellWidget(0, col, w)
        self.tabla_ingresos.setRowHeight(0, 40)

        for i, ing in enumerate(datos):
            row = i + 1
            item_id = QTableWidgetItem(ing['id_ingreso'])
            item_id.setData(Qt.UserRole, ing['id_ingreso'].lower())
            self.tabla_ingresos.setItem(row, 0, item_id)

            fecha = ing.get('fecha_ingreso')
            if fecha:
                item_fecha = QTableWidgetItem(fecha.strftime('%d/%m/%Y'))
                item_fecha.setData(Qt.UserRole, fecha)
            else:
                item_fecha = QTableWidgetItem('')
                item_fecha.setData(Qt.UserRole, _date.min)
            self.tabla_ingresos.setItem(row, 1, item_fecha)

            item_prov = QTableWidgetItem(ing.get('num_proveedor', ''))
            item_prov.setData(Qt.UserRole, ing.get('num_proveedor', '').lower())
            self.tabla_ingresos.setItem(row, 2, item_prov)

            item_rs = QTableWidgetItem(ing.get('razon_social', ''))
            item_rs.setData(Qt.UserRole, ing.get('razon_social', '').lower())
            self.tabla_ingresos.setItem(row, 3, item_rs)

            cant = ing.get('cantidad_remitos', 0)
            item_rem = QTableWidgetItem(str(cant))
            item_rem.setData(Qt.UserRole, cant)
            item_rem.setTextAlignment(Qt.AlignCenter)
            self.tabla_ingresos.setItem(row, 4, item_rem)

            nums_gen = self._nums_generados(ing['id_ingreso'])
            item_gen = QTableWidgetItem(nums_gen)
            item_gen.setData(Qt.UserRole, nums_gen.lower())
            self.tabla_ingresos.setItem(row, 5, item_gen)

        self.tabla_ingresos.blockSignals(False)
        header = self.tabla_ingresos.horizontalHeader()
        for c in range(self.tabla_ingresos.columnCount()):
            header.setSectionResizeMode(c, QHeaderView.Stretch)

    def _crear_filtro_fecha(self) -> QWidget:
        container = QWidget()
        lay = QHBoxLayout(container)
        lay.setContentsMargins(2, 2, 2, 2)
        lay.setSpacing(0)

        txt = QLineEdit()
        txt.setReadOnly(True)
        txt.setPlaceholderText('dd/mm/aaaa')
        if self._filtros[1]:
            txt.setText(self._filtros[1])

        btn_cal = QPushButton('\U0001F4C5')
        btn_cal.setFixedWidth(30)
        btn_cal.setToolTip('Abrir calendario')
        btn_cal.clicked.connect(lambda checked, t=txt: self._mostrar_calendario(t))

        btn_clear = QPushButton('\u2715')
        btn_clear.setFixedWidth(22)
        btn_clear.setToolTip('Limpiar filtro')
        btn_clear.clicked.connect(lambda checked, t=txt: self._limpiar_filtro_fecha(t))

        lay.addWidget(txt)
        lay.addWidget(btn_cal)
        lay.addWidget(btn_clear)
        return container

    def _mostrar_calendario(self, txt_fecha: QLineEdit):
        cal = QCalendarWidget()
        cal.setWindowFlags(Qt.Popup)
        cal.setGridVisible(True)
        if self._filtros[1]:
            try:
                d = _date(*map(int, reversed(self._filtros[1].split('/'))))
                cal.setSelectedDate(QDate(d.year, d.month, d.day))
            except (ValueError, TypeError):
                pass
        cal.setMinimumDate(QDate(2020, 1, 1))
        cal.dateSelected.connect(lambda qdate, t=txt_fecha: self._on_fecha_seleccionada(qdate, t))
        cal.exec_()

    def _on_fecha_seleccionada(self, qdate: QDate, txt_fecha: QLineEdit):
        d = qdate.toPyDate()
        date_str = d.strftime('%d/%m/%Y')
        txt_fecha.setText(date_str)
        self._filtros[1] = date_str
        self._aplicar_filtros_y_orden()

    def _limpiar_filtro_fecha(self, txt_fecha: QLineEdit):
        txt_fecha.clear()
        self._filtros[1] = ''
        self._aplicar_filtros_y_orden()

    def _crear_combo_filtro(self, attr: str, placeholder: str) -> QComboBox:
        cbo = QComboBox()
        cbo.addItem('')
        valores = sorted({i.get(attr, '') for i in self._ingresos if i.get(attr)})
        for v in valores:
            cbo.addItem(v)
        idx = cbo.findText(self._filtros.get(
            2 if attr == 'num_proveedor' else 3, ''))
        if idx >= 0:
            cbo.setCurrentIndex(idx)
        col = 2 if attr == 'num_proveedor' else 3
        cbo.currentTextChanged.connect(lambda text, c=col: self._on_filtro_changed(c, text))
        return cbo

    def _on_filtro_changed(self, col, text):
        self._filtros[col] = text
        self._aplicar_filtros_y_orden()

    def _limpiar_detalle(self):
        self._remitos = []
        self._remito_recepciones.clear()
        self._expanded_rows.clear()
        self._default_row_heights.clear()
        self.lbl_detalle_id.setText('')
        self.lbl_detalle_proveedor.setText('')
        self.lbl_remitos_generados.setText('')
        self.tabla_remitos.setRowCount(0)

    def _on_fecha_rango_changed(self):
        self._fecha_desde = self.dtp_desde.date() if self.dtp_desde.date().isValid() else None
        self._fecha_hasta = self.dtp_hasta.date() if self.dtp_hasta.date().isValid() else None
        self._aplicar_filtros_y_orden()

    def _limpiar_fechas(self):
        self.dtp_desde.setDate(QDate.currentDate().addMonths(-1))
        self.dtp_hasta.setDate(QDate.currentDate())
        self._fecha_desde = None
        self._fecha_hasta = None
        self._aplicar_filtros_y_orden()

    def _actualizar_resumen(self):
        total = len(self._ingresos_filtrados)
        total_bultos = sum(i.get('cantidad_remitos', 0) for i in self._ingresos_filtrados)
        win = self.window()
        if hasattr(win, 'statusBar'):
            win.statusBar().showMessage(
                f'Mostrando {total} ingresos | Filtro activo'
                if total != len(self._ingresos)
                else f'{total} ingresos cargados | {total_bultos} remitos'
            )

    def _exportar_excel(self):
        if not self._ingresos_filtrados:
            QMessageBox.information(self, 'Sin datos', 'No hay datos para exportar.')
            return
        path, _ = QFileDialog.getSaveFileName(
            self, 'Exportar a Excel', 'ingresos.xlsx',
            'Archivos Excel (*.xlsx)',
        )
        if not path:
            return
        try:
            QApplication.setOverrideCursor(QCursor(Qt.WaitCursor))
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, PatternFill
            from openpyxl.utils import get_column_letter
            wb = Workbook()
            ws = wb.active
            ws.title = 'Ingresos'

            encabezados = ['ID Ingreso', 'Fecha', 'Proveedor', 'Razón Social',
                           'Remitos', 'Remito Generado', 'Bultos', 'Recepción / Destino']
            e_enc = Font(bold=True, color='FFFFFF', size=10)
            r_enc = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')

            e_ing = Font(bold=True, size=10)
            r_ing = PatternFill(start_color='E8EAF6', end_color='E8EAF6', fill_type='solid')
            b_ing = Border(
                bottom=Side(style='medium', color='1565C0'),
                top=Side(style='thin', color='1565C0'),
            )

            e_rem = Font(bold=True, size=9, color='1565C0')
            r_rem = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
            b_rem = Border(bottom=Side(style='thin', color='BBDEFB'))

            e_dst = Font(bold=True, size=9, color='2E7D32')
            r_dst = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
            b_dst = Border(bottom=Side(style='dotted', color='C8E6C9'))

            e_rp = Font(size=9, color='666666')
            b_rp = Border(bottom=Side(style='hair', color='E0E0E0'))

            for col, h in enumerate(encabezados, 1):
                c = ws.cell(row=1, column=col, value=h)
                c.font = e_enc
                c.fill = r_enc
                c.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 22

            row = 2
            for ing in self._ingresos_filtrados:
                id_ing = ing['id_ingreso']
                fecha = ing.get('fecha_ingreso')
                remitos = self._svc.obtener_remitos_por_ingreso(id_ing)
                total_bultos = sum(r.get('cantidad_bultos', 0) for r in remitos)

                ws.cell(row=row, column=1, value=id_ing).font = e_ing
                ws.cell(row=row, column=2, value=fecha.strftime('%d/%m/%Y') if fecha else '').font = e_ing
                ws.cell(row=row, column=3, value=ing.get('num_proveedor', '')).font = e_ing
                ws.cell(row=row, column=4, value=ing.get('razon_social', '')).font = e_ing
                ws.cell(row=row, column=5, value=len(remitos)).font = e_ing
                ws.cell(row=row, column=5).alignment = Alignment(horizontal='center')
                ws.cell(row=row, column=6, value=self._nums_generados(id_ing)).font = e_ing
                ws.cell(row=row, column=7, value=total_bultos).font = e_ing
                ws.cell(row=row, column=7).alignment = Alignment(horizontal='center')
                remitos_nums = ', '.join(r.get('num_remito', '') for r in remitos)
                ws.cell(row=row, column=8, value=f'Remitos: {remitos_nums}').font = e_ing
                for col in range(1, 9):
                    ws.cell(row=row, column=col).fill = r_ing
                    ws.cell(row=row, column=col).border = b_ing
                row += 1

                destinos_global: dict[str, dict] = {}
                for r in remitos:
                    recepciones = self._svc.obtener_recepciones_por_remito(r['id_remito'])
                    for rp in recepciones:
                        suc = rp.get('sucursal_destino', 'S/D')
                        if suc not in destinos_global:
                            destinos_global[suc] = {
                                'nombre': rp.get('nombre_destino', suc),
                                'domicilio': rp.get('domicilio_cliente', ''),
                                'ciudad': rp.get('ciudad_cliente', ''),
                                'rps': [],
                                'remitos_set': set(),
                                'bultos': 0,
                            }
                        destinos_global[suc]['rps'].append({
                            'id': rp.get('id_recepcion_prevista', ''),
                            'destino': rp.get('sucursal_destino', 'S/D'),
                            'nombre_destino': rp.get('nombre_destino', ''),
                            'domicilio': rp.get('domicilio_cliente', ''),
                        })
                        destinos_global[suc]['remitos_set'].add(r.get('num_remito', ''))
                        destinos_global[suc]['bultos'] += r.get('cantidad_bultos', 0)

                self._enrich_destinos_from_sqlserver(destinos_global)

                for suc, info in destinos_global.items():
                    nombre = info['nombre']
                    dom = info['domicilio']
                    ciudad = info['ciudad']
                    cant_rps = len(info['rps'])
                    cant_rem = len(info['remitos_set'])
                    destino_txt = f'{suc} : {nombre}' if nombre and nombre != suc else suc
                    if dom:
                        destino_txt += f' ({dom})'
                    detalle_parts = []
                    if ciudad:
                        detalle_parts.append(ciudad)
                    detalle_parts.append(f'{cant_rps} recepcion(es)')
                    if cant_rem > 1:
                        detalle_parts.append(f'en {cant_rem} remitos')
                    detalle_parts.append(f'{info["bultos"]} bultos')
                    detalle_txt = ' — '.join(detalle_parts)

                    ws.cell(row=row, column=8, value=f'▸ {destino_txt}').font = e_dst
                    for col in range(1, 9):
                        ws.cell(row=row, column=col).fill = r_dst
                        ws.cell(row=row, column=col).border = b_dst
                    row += 1

                    ws.cell(row=row, column=8, value=f'   {detalle_txt}').font = e_rp
                    for col in range(1, 9):
                        ws.cell(row=row, column=col).border = b_rp
                    row += 1

                    for rp_info in info['rps']:
                        rp_id = rp_info['id']
                        rp_linea = f'     • {rp_id}'
                        ws.cell(row=row, column=8, value=rp_linea).font = e_rp
                        for col in range(1, 9):
                            ws.cell(row=row, column=col).border = b_rp
                        row += 1

                row += 1

            ws.column_dimensions['A'].width = 22
            ws.column_dimensions['B'].width = 13
            ws.column_dimensions['C'].width = 14
            ws.column_dimensions['D'].width = 32
            ws.column_dimensions['E'].width = 10
            ws.column_dimensions['F'].width = 20
            ws.column_dimensions['G'].width = 9
            ws.column_dimensions['H'].width = 65
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = 'A2'
            wb.save(path)
            QApplication.restoreOverrideCursor()
            QMessageBox.information(
                self, 'Exportado',
                f'Se exportaron {len(self._ingresos_filtrados)} ingresos a:\n{path}'
            )
        except Exception as e:
            QApplication.restoreOverrideCursor()
            QMessageBox.critical(self, 'Error', f'Error al exportar:\n{e}')

    def _on_seleccion_cambiada(self):
        rows = self.tabla_ingresos.selectedIndexes()
        if not rows:
            self._limpiar_detalle()
            self.btn_eliminar.setEnabled(False)
            self.btn_reimprimir_remitos.setEnabled(False)
            self.btn_reimprimir_etiquetas.setEnabled(False)
            return
        row = rows[0].row()
        if row < 1 or row - 1 >= len(self._ingresos_filtrados):
            self.btn_eliminar.setEnabled(False)
            self.btn_reimprimir_remitos.setEnabled(False)
            self.btn_reimprimir_etiquetas.setEnabled(False)
            return
        self.btn_eliminar.setEnabled(True)
        self.btn_reimprimir_remitos.setEnabled(bool(self._imp_svc))
        self.btn_reimprimir_etiquetas.setEnabled(bool(self._imp_svc))
        ingreso = self._ingresos_filtrados[row - 1]
        self._cargar_detalle(ingreso['id_ingreso'])

    def _cargar_detalle(self, id_ingreso: str):
        try:
            remitos = self._svc.obtener_remitos_por_ingreso(id_ingreso)
            self._remitos = remitos
            self._remito_recepciones.clear()
            for r in remitos:
                recepciones = self._svc.obtener_recepciones_por_remito(r['id_remito'])
                self._remito_recepciones[r['id_remito']] = recepciones

            ingreso = next((i for i in self._ingresos if i['id_ingreso'] == id_ingreso), None)
            if ingreso:
                self.lbl_detalle_id.setText(
                    f'ID: {ingreso["id_ingreso"]} — Fecha: {ingreso["fecha_ingreso"].strftime("%d/%m/%Y")}'
                )
                self.lbl_detalle_proveedor.setText(
                    f'Proveedor: {ingreso.get("num_proveedor", "")} - {ingreso.get("razon_social", "")}'
                )
            generados = self._svc.obtener_remitos_generados(id_ingreso)
            if generados:
                partes = []
                for g in generados:
                    num = g.get('num_remito') or ''
                    est = self.ESTADOS_REMITO.get(int(g.get('estado') or 1), 'Emitido')
                    partes.append(f'{num} ({est})' if num else est)
                self.lbl_remitos_generados.setText(
                    'Remitos generados: ' + ', '.join(partes)
                )
            else:
                self.lbl_remitos_generados.setText('')
            self._expanded_rows.clear()
            self._default_row_heights.clear()
            self._poblar_tabla_remitos()
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al cargar detalle:\n{e}')

    def _poblar_tabla_remitos(self):
        self.tabla_remitos.setRowCount(len(self._remitos))
        for i, r in enumerate(self._remitos):
            self.tabla_remitos.setItem(i, 0, QTableWidgetItem(r['num_remito']))
            fecha = r.get('fecha_remito')
            if fecha:
                self.tabla_remitos.setItem(i, 1, QTableWidgetItem(fecha.strftime('%d/%m/%Y')))
            else:
                self.tabla_remitos.setItem(i, 1, QTableWidgetItem(''))
            self.tabla_remitos.setItem(i, 2, QTableWidgetItem(str(r.get('cantidad_bultos', 0))))
            self.tabla_remitos.setItem(i, 3, QTableWidgetItem(str(r.get('cantidad_unidades', 0))))
            unidades = r.get('cantidad_unidades', 0)
            valor_unit = r.get('valor_declarado_unitario', 0)
            total_vd = unidades * valor_unit if unidades > 0 and valor_unit > 0 else 0
            item_vd = QTableWidgetItem(f'${total_vd:,.2f}' if total_vd > 0 else '-')
            item_vd.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.tabla_remitos.setItem(i, 4, item_vd)
            receps = self._remito_recepciones.get(r['id_remito'], [])
            texto = ', '.join(
                f'{rep["id_recepcion_prevista"]} ({rep.get("sucursal_destino", "")})'
                for rep in receps
            )
            self.tabla_remitos.setItem(i, 5, QTableWidgetItem(texto if texto else '—'))
        header = self.tabla_remitos.horizontalHeader()
        for c in range(self.tabla_remitos.columnCount()):
            if c == 5:
                header.setSectionResizeMode(c, QHeaderView.Stretch)
            else:
                header.setSectionResizeMode(c, QHeaderView.ResizeToContents)

    def _on_remito_cell_clicked(self, row, col):
        if col != self.COL_RECEPCIONES:
            return
        if row < 0 or row >= len(self._remitos):
            return
        if row in self._expanded_rows:
            altura_orig = self._default_row_heights.pop(row, 25)
            self.tabla_remitos.setRowHeight(row, altura_orig)
            self._expanded_rows.discard(row)
        else:
            if row not in self._default_row_heights:
                self._default_row_heights[row] = self.tabla_remitos.rowHeight(row)
            item = self.tabla_remitos.item(row, self.COL_RECEPCIONES)
            if not item:
                return
            ancho = self.tabla_remitos.columnWidth(self.COL_RECEPCIONES) - 10
            fm = QFontMetrics(self.tabla_remitos.font())
            rect = fm.boundingRect(0, 0, ancho, 0, Qt.AlignLeft | Qt.TextWordWrap, item.text())
            self.tabla_remitos.setRowHeight(row, rect.height() + 10)
            self._expanded_rows.add(row)

    def _eliminar_ingreso(self):
        filas = sorted({
            idx.row() for idx in self.tabla_ingresos.selectedIndexes()
            if 0 <= idx.row() - 1 < len(self._ingresos_filtrados)
        })
        if not filas:
            return
        ingresos = [self._ingresos_filtrados[r - 1] for r in filas]

        if len(ingresos) == 1:
            ing = ingresos[0]
            detalle = (
                f'ID: {ing["id_ingreso"]}\n'
                f'Proveedor: {ing.get("razon_social", ing.get("num_proveedor", ""))}'
            )
        else:
            lineas = '\n'.join(
                f'• {i["id_ingreso"]} - {i.get("razon_social", i.get("num_proveedor", ""))}'
                for i in ingresos
            )
            detalle = f'Se eliminarán los siguientes {len(ingresos)} ingresos:\n\n{lineas}'

        reply = QMessageBox.question(
            self, 'Confirmar Eliminación',
            f'¿Está seguro de eliminar el/los ingreso(s)?\n\n'
            f'{detalle}\n\n'
            f'Esta acción no se puede deshacer.',
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if reply != QMessageBox.Yes:
            return

        login_dlg = LoginDialog(self._sqlserver_config, self)
        if login_dlg.exec_() != LoginDialog.Accepted:
            return

        palabra_clave, ok = QInputDialog.getText(
            self, 'Validación de Seguridad',
            'Ingrese la palabra clave para confirmar la eliminación:',
            QLineEdit.Password
        )
        if not ok:
            return
        if palabra_clave != self._clave_anulacion:
            QMessageBox.warning(
                self, 'Acceso Denegado',
                'La palabra clave es incorrecta. No se realizó ninguna eliminación.'
            )
            return

        eliminados = []
        errores = []
        for ing in ingresos:
            try:
                self._svc.eliminar_ingreso(ing['id_ingreso'])
                eliminados.append(ing['id_ingreso'])
            except Exception as e:
                errores.append(f'{ing["id_ingreso"]}: {e}')

        if eliminados:
            QMessageBox.information(
                self, 'Eliminados',
                'Se eliminaron exitosamente:\n' + '\n'.join(eliminados)
            )
        if errores:
            QMessageBox.critical(
                self, 'Error',
                'No se pudieron eliminar:\n' + '\n'.join(errores)
            )
        self._cargar_ingresos()

    def _reimprimir(self, tipo: str):
        if not self._imp_svc:
            QMessageBox.warning(self, 'Sin servicio', 'No hay servicio de impresión disponible.')
            return
        rows = self.tabla_ingresos.selectedIndexes()
        if not rows:
            return
        row = rows[0].row()
        if row < 1 or row - 1 >= len(self._ingresos_filtrados):
            return
        id_ingreso = self._ingresos_filtrados[row - 1]['id_ingreso']

        try:
            ingreso = self._svc.obtener_ingreso_completo(id_ingreso)
            if not ingreso:
                QMessageBox.warning(self, 'Error', 'No se pudo cargar el ingreso.')
                return
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al cargar ingreso para reimprimir:\n{e}')
            return

        impresora = self._seleccionar_impresora(tipo)
        if impresora is None:
            return

        try:
            if tipo == 'remitos':
                exitos, errores = self._imp_svc.imprimir_remitos(
                    ingreso, impresora, reimprimir=True
                )
            else:
                exitos, errores = self._imp_svc.imprimir_etiquetas(
                    ingreso, impresora, reimprimir=True
                )
            if errores:
                QMessageBox.warning(
                    self, 'Reimpresión',
                    'No se pudieron generar:\n' + '\n'.join(errores)
                )
            elif exitos:
                QMessageBox.information(
                    self, 'Reimpresión',
                    'Documentos generados:\n' + '\n'.join(exitos)
                )
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al reimprimir:\n{e}')

    def _seleccionar_impresora(self, tipo: str) -> str | None:
        try:
            import win32print
            impresoras = list(win32print.EnumPrinters(
                win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
            ))
            nombres = sorted(set(p[2] for p in impresoras))
        except Exception:
            nombres = []

        if not nombres:
            QMessageBox.information(
                self, 'Sin impresoras',
                'No se encontraron impresoras disponibles.\n'
                'Se abrirá el PDF para visualización.'
            )
            return ''

        cfg_actual = ''
        if self._imp_svc:
            if tipo == 'remitos':
                cfg_actual = self._imp_svc._cm.impresoras.impresora_remitos
            else:
                cfg_actual = self._imp_svc._cm.impresoras.impresora_etiquetas

        idx_default = 0
        if cfg_actual in nombres:
            idx_default = nombres.index(cfg_actual)

        nombre, ok = QInputDialog.getItem(
            self, f'Seleccionar impresora - {tipo.title()}',
            f'Impresora para {tipo}:',
            nombres, idx_default, False
        )
        if not ok:
            return None
        return nombre

    def _enrich_destinos_from_sqlserver(self, destinos: dict):
        if not self._sqlserver_config:
            return
        ids_recepciones = set()
        for suc, info in destinos.items():
            for rp in info['rps']:
                if not info['nombre'] or info['nombre'] == suc or not info['domicilio']:
                    ids_recepciones.add(rp['id'])
        if not ids_recepciones:
            return
        try:
            ss_conn = SQLServerConnection(self._sqlserver_config)
            ss_conn.connect()
            ss_rows = ss_conn.buscar_recepciones_por_ids(list(ids_recepciones))
            ss_conn.disconnect()
            by_id = {str(r['IDRecepcion']): r for r in ss_rows}
            for suc, info in destinos.items():
                if info['nombre'] and info['nombre'] != suc and info['domicilio']:
                    continue
                for rp in info['rps']:
                    ss_data = by_id.get(str(rp['id']), {})
                    if ss_data:
                        if not info['nombre'] or info['nombre'] == suc:
                            info['nombre'] = ss_data.get('NombreDestino', '') or info['nombre']
                        if not info['domicilio']:
                            info['domicilio'] = ss_data.get('DomicilioCliente', '')
                        if not info['ciudad']:
                            info['ciudad'] = ss_data.get('CiudadCliente', '')
                        rp['nombre_destino'] = info['nombre']
                        rp['domicilio'] = info['domicilio']
        except Exception:
            pass
